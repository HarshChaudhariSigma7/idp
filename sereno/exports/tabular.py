"""CSV / Excel exports shaped for ERP import. One flat row per line item (header fields repeated),
which is what Tally / SAP Business One / regional-ERP import utilities expect. Column names and
date format can be remapped per customer (tenant.settings.export_profiles) without code; a native
ERP connector is built only after the first customer confirms which ERP they run."""
from __future__ import annotations

import csv
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from sereno.extraction.doc_specs import SPECS
from sereno.review import effective_value


def _fmt(v, vtype: str, date_fmt: str):
    if v is None:
        return ""
    if vtype == "date":
        try:
            return date.fromisoformat(v).strftime(date_fmt)
        except ValueError:
            return v
    return v


def rows_for(docs, doc_type: str, profile: dict | None = None, layout: str = "lines") -> tuple[list[str], list[list]]:
    spec = SPECS[doc_type]
    profile = profile or {}
    date_fmt = profile.get("date_format", "%d-%m-%Y")
    rename = {c["source"]: c["header"] for c in profile.get("columns", []) if c.get("source") and c.get("header")}
    head_names = [f.name for f in spec.fields]
    line_names = [f.name for f in spec.line_fields] if layout == "lines" else []
    labels = {f.name: f.label for f in spec.fields}
    cols = ["Document ID", "Status"] + [rename.get(n, labels[n]) for n in head_names]
    if line_names:
        llabels = {f.name: f.label for f in spec.line_fields}
        cols += ["Line no."] + [rename.get(f"line.{n}", f"Line {llabels[n]}") for n in line_names]
    types = {f.name: f.type for f in spec.fields}
    ltypes = {f.name: f.type for f in spec.line_fields}
    out = []
    for d in docs:
        header_vals, lines = {}, {}
        for f in d.fields:
            v = effective_value(f)
            if f.line_index is None:
                header_vals[f.field_name] = v
            else:
                lines.setdefault(f.line_index, {})[f.field_name] = v
        base = [d.id, "Reviewed" if d.reviewed_at else "Auto-verified"] + \
               [_fmt(header_vals.get(n), types[n], date_fmt) for n in head_names]
        if line_names and lines:
            for i in sorted(lines):
                out.append(base + [i + 1] + [_fmt(lines[i].get(n), ltypes[n], date_fmt) for n in line_names])
        else:
            out.append(base + ([""] * (len(line_names) + 1) if line_names else []))
    return cols, out


def to_csv(cols: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(cols)
    for r in rows:
        w.writerow([_safe_cell(c) for c in r])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens Hindi text correctly


def _safe_cell(v):
    """Neutralise spreadsheet formula injection from document text."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") and not _is_number(v):
        return "'" + v
    return v


def _is_number(s: str) -> bool:
    try:
        float(s)
        return True
    except ValueError:
        return False


def to_xlsx(sheets: dict[str, tuple[list[str], list[list]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, (cols, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
        for r in rows:
            ws.append([_safe_cell(c) for c in r])
        for i, col in enumerate(cols, 1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r)]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
